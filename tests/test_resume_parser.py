"""Real PDF/DOCX extraction fixtures and offline agent CLI integration."""

import hashlib
import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from docx.enum.section import WD_SECTION
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject, DictionaryObject, NameObject, NumberObject,
)

from src.resume.parser import (
    EncryptedPDFError, NoResumeTextError, OCRRequiredError,
    ResumeExtractionError, extract_resume_text,
)


def write_pdf(path, pages, password=None):
    """Write real small PDFs: each page is text, blank, or an image-only scan."""
    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        content = DecodedStreamObject()
        if text == "<scan>":
            image = DecodedStreamObject()
            image.set_data(b"\xff\xff\xff")
            image.update({
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(1),
                NameObject("/Height"): NumberObject(1),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/BitsPerComponent"): NumberObject(8),
            })
            # Indirect resources are common in exported PDFs.
            objects = writer._add_object(DictionaryObject({NameObject("/Im"): writer._add_object(image)}))
            resources = DictionaryObject({NameObject("/XObject"): objects})
            page[NameObject("/Resources")] = writer._add_object(resources)
            content.set_data(b"q 250 0 0 250 0 0 cm /Im Do Q")
        elif text:
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
            })
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content.set_data(f"BT /F1 12 Tf 20 250 Td ({escaped}) Tj ET".encode("ascii"))
        else:
            continue
        page[NameObject("/Contents")] = writer._add_object(content)
    if password is not None:
        writer.encrypt(password)
    writer.write(path)


class ResumeParserTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def test_pdf_reads_all_pages_in_order_and_skips_blank_pages(self):
        path = self.root / "resume.PDF"
        write_pdf(path, ["Built Python pipelines.", None, "Created SQL reports."])
        self.assertEqual(extract_resume_text(path), "Built Python pipelines.\n\nCreated SQL reports.")

    def test_image_only_and_mixed_scanned_pdfs_need_ocr(self):
        for pages, number in [(["<scan>"], "1"), (["Python skills.", "<scan>"], "2")]:
            path = self.root / "scan.pdf"
            write_pdf(path, pages)
            with self.subTest(pages=pages), self.assertRaisesRegex(OCRRequiredError, f"page\\(s\\) {number}"):
                extract_resume_text(path)

    def test_blank_pdf_and_empty_docx_raise_no_text(self):
        pdf, docx = self.root / "empty.pdf", self.root / "empty.docx"
        write_pdf(pdf, [None])
        Document().save(docx)
        for path in (pdf, docx):
            with self.subTest(path=path), self.assertRaises(NoResumeTextError):
                extract_resume_text(path)

    def test_password_protected_pdf_requires_correct_password(self):
        path = self.root / "encrypted.pdf"
        write_pdf(path, ["Python skills."], password="resume-password")
        for password in (None, "wrong-password"):
            with self.subTest(password=password), self.assertRaises(EncryptedPDFError):
                extract_resume_text(path, password=password)
        self.assertEqual(extract_resume_text(path, password="resume-password"), "Python skills.")

    def test_docx_keeps_paragraph_table_and_nested_table_order(self):
        path = self.root / "resume.DOCX"
        document = Document()
        document.add_paragraph("Experience")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "SQL"
        nested = table.cell(0, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Power BI"
        document.add_paragraph("Education: Computer Science")
        document.save(path)
        self.assertEqual(
            extract_resume_text(path),
            "Experience\nPython\tSQL\nPower BI\nEducation: Computer Science",
        )

    def test_docx_headers_footers_and_merged_cells_are_not_duplicated(self):
        path = self.root / "layout.docx"
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Candidate: Aisha"
        document.sections[0].footer.paragraphs[0].text = "Contact: aisha@example.test"
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Python SQL"
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("Second section")
        document.save(path)
        text = extract_resume_text(path)
        for expected in ("Candidate: Aisha", "Python SQL", "Contact: aisha@example.test"):
            self.assertEqual(text.count(expected), 1)
        self.assertLess(text.index("Candidate: Aisha"), text.index("Python SQL"))

    def test_docx_active_first_and_even_page_headers_are_included(self):
        path = self.root / "headers.docx"
        document = Document()
        section = document.sections[0]
        section.first_page_header.paragraphs[0].text = "First-page contact"
        section.even_page_header.paragraphs[0].text = "Even-page contact"
        document.add_paragraph("Python")
        document.save(path)
        self.assertEqual(extract_resume_text(path), "Python")
        section.different_first_page_header_footer = True
        document.settings.odd_and_even_pages_header_footer = True
        document.save(path)
        text = extract_resume_text(path)
        self.assertIn("First-page contact", text)
        self.assertIn("Even-page contact", text)

    def test_utf8_bom_unicode_line_breaks_and_empty_text(self):
        path = self.root / "resume.txt"
        path.write_bytes("\ufeffZoë 李\r\nPython SQL\r\n".encode("utf-8"))
        self.assertEqual(extract_resume_text(path), "Zoë 李\nPython SQL")
        path.write_text(" \n\t", encoding="utf-8")
        with self.assertRaises(NoResumeTextError):
            extract_resume_text(path)
        path.write_bytes(b"\xff\xfe\x00")
        with self.assertRaisesRegex(ResumeExtractionError, "UTF-8"):
            extract_resume_text(path)

    def test_missing_directories_unsupported_and_corrupt_files(self):
        with self.assertRaises(FileNotFoundError):
            extract_resume_text(self.root / "missing.pdf")
        folder = self.root / "folder.pdf"
        folder.mkdir()
        with self.assertRaises(IsADirectoryError):
            extract_resume_text(folder)
        for suffix in (".doc", ".png"):
            with self.subTest(suffix=suffix), self.assertRaisesRegex(ResumeExtractionError, "Unsupported"):
                extract_resume_text(self.root / ("resume" + suffix))
        for suffix in (".pdf", ".docx"):
            path = self.root / ("broken" + suffix)
            path.write_bytes(b"This is not a document")
            with self.subTest(suffix=suffix), self.assertRaisesRegex(ResumeExtractionError, "Could not read"):
                extract_resume_text(path)

    def test_extraction_does_not_change_original_files(self):
        pdf, docx = self.root / "resume.pdf", self.root / "resume.docx"
        write_pdf(pdf, ["Python"])
        document = Document()
        document.add_paragraph("Python")
        document.save(docx)
        for path in (pdf, docx):
            before = hashlib.sha256(path.read_bytes()).digest()
            extract_resume_text(path)
            self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_agent_cli_extracts_pdf_and_docx_before_assessment(self):
        from src.models.model import main

        job = self.root / "job.txt"
        job.write_text("Python required", encoding="utf-8")
        pdf, docx = self.root / "resume.pdf", self.root / "resume.docx"
        write_pdf(pdf, ["Built Python pipelines."])
        document = Document()
        document.add_paragraph("Built Python pipelines.")
        document.save(docx)
        for path in (pdf, docx):
            agent = Mock()
            agent.assess_resume.return_value.model_dump_json.return_value = '{"fit": "strong_fit"}'
            with self.subTest(path=path), patch("src.models.model.JobMatchAgent", return_value=agent), \
                    patch.object(sys, "argv", ["agent", "--job", str(job), "--resume", str(path)]), \
                    redirect_stdout(io.StringIO()):
                self.assertEqual(main(), 0)
            agent.assess_resume.assert_called_once_with("Python required", "Built Python pipelines.")

    def test_scanned_pdf_cli_fails_before_constructing_agent(self):
        from src.models.model import main

        job, resume = self.root / "job.txt", self.root / "scan.pdf"
        job.write_text("Python required", encoding="utf-8")
        write_pdf(resume, ["<scan>"])
        with patch("src.models.model.JobMatchAgent") as agent, \
                patch.object(sys, "argv", ["agent", "--job", str(job), "--resume", str(resume)]), \
                patch("sys.stderr", new=io.StringIO()) as error_stream:
            with self.assertRaises(SystemExit) as error:
                main()
            self.assertEqual(error.exception.code, 1)
            self.assertIn("OCR", error_stream.getvalue())
            agent.assert_not_called()


if __name__ == "__main__":
    unittest.main()
